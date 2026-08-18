from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook

from sleep_ai_scientist.feature_extraction.extractors.scale_workbook import (
    extract_scale_workbooks,
    select_baseline_scale_rows,
)


SCALE_HEADERS = ["PSQI", "ISI", "嗜睡", "BAI", "BDI"]


def _write_scale_workbook(root: Path, subject: str, visits: list[tuple[str, list[object]]]) -> Path:
    target = root / subject / "scales" / f"{subject}_scales.xlsx"
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "scales"
    labels = {
        "0w": "id基线（0w）",
        "2w": "id治疗后评估（2w）",
        "3w": "id随访（3w）",
        "6w": "id随访（6w）",
    }
    for visit, values in visits:
        sheet.append([labels[visit], *SCALE_HEADERS])
        sheet.append([subject.replace("sub-", ""), *values])
        sheet.append([])
    workbook.save(target)
    return target


def test_extract_scale_workbooks_normalizes_visits_and_scores(tmp_path: Path) -> None:
    source = _write_scale_workbook(
        tmp_path,
        "sub-ISM035",
        [("0w", [13, 18, 1, 23, 22]), ("2w", [10, 15, 3, 15, 20])],
    )

    long_frame, audit = extract_scale_workbooks(tmp_path)

    assert list(long_frame.columns) == [
        "subject_id",
        "subject",
        "visit",
        "ISI",
        "PSQI",
        "BAI",
        "BDI",
        "sleepiness",
        "source_file",
    ]
    baseline = long_frame.set_index(["subject_id", "visit"]).loc[("sub-ISM035", "0w")]
    assert baseline["ISI"] == 18
    assert baseline["PSQI"] == 13
    assert baseline["BAI"] == 23
    assert baseline["BDI"] == 22
    assert baseline["sleepiness"] == 1
    assert baseline["source_file"] == str(source)
    assert audit["discovered_file_count"] == 1
    assert audit["parsed_file_count"] == 1
    assert audit["visit_counts"] == {"0w": 1, "2w": 1}


def test_extract_scale_workbooks_preserves_missing_scores_and_skips_corrupt_files(tmp_path: Path) -> None:
    _write_scale_workbook(tmp_path, "sub-YZHC001", [("0w", [4, None, 0, 2, 1])])
    corrupt = tmp_path / "sub-BROKEN" / "scales" / "sub-BROKEN_scales.xlsx"
    corrupt.parent.mkdir(parents=True)
    corrupt.write_text("not an xlsx archive", encoding="utf-8")

    long_frame, audit = extract_scale_workbooks(tmp_path)

    row = long_frame.iloc[0]
    assert pd.isna(row["ISI"])
    assert row["PSQI"] == 4
    assert audit["discovered_file_count"] == 2
    assert audit["parsed_file_count"] == 1
    assert audit["skipped_file_count"] == 1
    assert audit["skipped_files"][0]["path"] == str(corrupt)


def test_extract_scale_workbooks_audits_readable_files_without_data_rows(tmp_path: Path) -> None:
    target = tmp_path / "sub-EMPTY001" / "scales" / "sub-EMPTY001_scales.xlsx"
    target.parent.mkdir(parents=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["id基线（0w）", *SCALE_HEADERS])
    sheet.append([])
    workbook.save(target)

    long_frame, audit = extract_scale_workbooks(tmp_path)

    assert long_frame.empty
    assert audit["parsed_file_count"] == 1
    assert audit["empty_file_count"] == 1
    assert audit["empty_files"] == [str(target)]
    assert audit["subjects_without_usable_visit"] == ["sub-EMPTY001"]


def test_extract_scale_workbooks_keeps_first_non_missing_duplicate_value(tmp_path: Path) -> None:
    source = _write_scale_workbook(
        tmp_path,
        "sub-DUP001",
        [("0w", [8, None, 1, 4, 5]), ("0w", [9, 12, 2, 6, 7])],
    )

    long_frame, audit = extract_scale_workbooks(tmp_path)

    assert len(long_frame) == 1
    row = long_frame.iloc[0]
    assert row["PSQI"] == 8
    assert row["ISI"] == 12
    assert row["sleepiness"] == 1
    assert row["source_file"] == str(source)
    assert audit["duplicate_subject_visit_count"] == 1


def test_extract_scale_workbooks_does_not_assign_unsupported_followups_to_previous_visit(tmp_path: Path) -> None:
    target = tmp_path / "sub-LONG001" / "scales" / "sub-LONG001_scales.xlsx"
    target.parent.mkdir(parents=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["id随访（6w）", *SCALE_HEADERS])
    sheet.append(["LONG001", 6, 7, 1, 2, 3])
    sheet.append([])
    sheet.append(["id随访（10w）", *SCALE_HEADERS])
    sheet.append(["LONG001", 4, 5, 0, 1, 2])
    workbook.save(target)

    long_frame, audit = extract_scale_workbooks(tmp_path)

    assert list(long_frame["visit"]) == ["6w"]
    assert long_frame.iloc[0]["ISI"] == 7
    assert audit["duplicate_subject_visit_count"] == 0


def test_select_baseline_scale_rows_prefers_observed_zero_week() -> None:
    long_frame = pd.DataFrame(
        [
            _scale_row("sub-BASE", "2w", isi=11),
            _scale_row("sub-BASE", "0w", isi=18),
        ]
    )

    baseline, audit = select_baseline_scale_rows(long_frame)

    row = baseline.iloc[0]
    assert row["ISI"] == 18
    assert row["baseline_source_visit"] == "0w"
    assert bool(row["baseline_is_substituted"]) is False
    assert audit["substituted_baseline_count"] == 0


def test_select_baseline_scale_rows_uses_closest_available_followup() -> None:
    long_frame = pd.DataFrame(
        [
            _scale_row("sub-TWO", "6w", isi=6),
            _scale_row("sub-TWO", "2w", isi=12),
            _scale_row("sub-TWO", "3w", isi=9),
            _scale_row("sub-THREE", "6w", isi=7),
            _scale_row("sub-THREE", "3w", isi=10),
        ]
    )

    baseline, audit = select_baseline_scale_rows(long_frame)

    indexed = baseline.set_index("subject_id")
    assert indexed.loc["sub-TWO", "ISI"] == 12
    assert indexed.loc["sub-TWO", "baseline_source_visit"] == "2w"
    assert indexed.loc["sub-THREE", "ISI"] == 10
    assert indexed.loc["sub-THREE", "baseline_source_visit"] == "3w"
    assert indexed["baseline_is_substituted"].astype(bool).all()
    assert audit["substituted_baseline_count"] == 2


def _scale_row(subject: str, visit: str, *, isi: float) -> dict[str, object]:
    return {
        "subject_id": subject,
        "subject": subject,
        "visit": visit,
        "ISI": isi,
        "PSQI": 5,
        "BAI": 4,
        "BDI": 3,
        "sleepiness": 2,
        "source_file": f"/{subject}_scales.xlsx",
    }
