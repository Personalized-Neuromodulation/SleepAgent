from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook


SCORE_COLUMNS = ["ISI", "PSQI", "BAI", "BDI", "sleepiness"]
LONG_COLUMNS = ["subject_id", "subject", "visit", *SCORE_COLUMNS, "source_file"]
BASELINE_COLUMNS = [
    "subject_id",
    "subject",
    *SCORE_COLUMNS,
    "baseline_source_visit",
    "baseline_is_substituted",
]
VISIT_PRIORITY = {"0w": 0, "2w": 1, "3w": 2, "6w": 3}
VISIT_PATTERN = re.compile(r"(?<!\d)(\d+)\s*w", re.IGNORECASE)


def extract_scale_workbooks(root: str | Path) -> tuple[pd.DataFrame, dict[str, Any]]:
    root_path = Path(root)
    paths = sorted(path for path in root_path.rglob("*_scales.xlsx") if "scales" in path.parts)
    rows: list[dict[str, Any]] = []
    skipped_files: list[dict[str, str]] = []
    empty_files: list[str] = []
    subjects_without_usable_visit: list[str] = []
    parsed_file_count = 0

    for path in paths:
        subject = _subject_from_path(path)
        if not subject:
            skipped_files.append({"path": str(path), "error": "No sub-* identifier in path."})
            continue
        try:
            workbook_rows = _parse_workbook(path, subject)
            rows.extend(workbook_rows)
            parsed_file_count += 1
            if not workbook_rows:
                empty_files.append(str(path))
                subjects_without_usable_visit.append(subject)
        except Exception as exc:
            skipped_files.append({"path": str(path), "error": f"{type(exc).__name__}: {exc}"})

    frame = pd.DataFrame(rows, columns=LONG_COLUMNS)
    frame, duplicate_count = _combine_duplicate_visits(frame)
    visit_counts = {
        str(visit): int(count)
        for visit, count in frame["visit"].value_counts(sort=False).items()
    } if not frame.empty else {}
    audit = {
        "source_root": str(root_path),
        "discovered_file_count": len(paths),
        "parsed_file_count": parsed_file_count,
        "skipped_file_count": len(skipped_files),
        "skipped_files": skipped_files,
        "empty_file_count": len(empty_files),
        "empty_files": empty_files,
        "subjects_without_usable_visit": subjects_without_usable_visit,
        "duplicate_subject_visit_count": duplicate_count,
        "visit_counts": visit_counts,
    }
    return frame, audit


def select_baseline_scale_rows(long_frame: pd.DataFrame) -> tuple[pd.DataFrame, dict[str, Any]]:
    required = {"subject_id", "subject", "visit", *SCORE_COLUMNS}
    missing = sorted(required - set(long_frame.columns))
    if missing:
        raise ValueError(f"Scale long table is missing required columns: {missing}")

    candidates = long_frame[long_frame["visit"].isin(VISIT_PRIORITY)].copy()
    if candidates.empty:
        return pd.DataFrame(columns=BASELINE_COLUMNS), {
            "baseline_subject_count": 0,
            "substituted_baseline_count": 0,
        }
    candidates["_visit_rank"] = candidates["visit"].map(VISIT_PRIORITY)
    candidates = candidates.sort_values(["subject_id", "_visit_rank"], kind="stable")
    baseline = candidates.drop_duplicates("subject_id", keep="first").copy()
    baseline = baseline.rename(columns={"visit": "baseline_source_visit"})
    baseline["baseline_is_substituted"] = baseline["baseline_source_visit"].ne("0w")
    baseline = baseline[BASELINE_COLUMNS].sort_values("subject_id").reset_index(drop=True)
    audit = {
        "baseline_subject_count": int(len(baseline)),
        "substituted_baseline_count": int(baseline["baseline_is_substituted"].sum()),
    }
    return baseline, audit


def _parse_workbook(path: Path, subject: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows: list[dict[str, Any]] = []
    try:
        for sheet in workbook.worksheets:
            current_visit = ""
            column_map: dict[str, int] = {}
            for values in sheet.iter_rows(values_only=True):
                is_visit_header, visit = _visit_header_from_row(values)
                if is_visit_header:
                    current_visit = visit
                    column_map = _score_column_map(values) if visit else {}
                    continue
                if not current_visit or not column_map or not any(value is not None for value in values):
                    continue
                row = {
                    "subject_id": subject,
                    "subject": subject,
                    "visit": current_visit,
                    **{
                        score: _numeric_value(values[index] if index < len(values) else None)
                        for score, index in column_map.items()
                    },
                    "source_file": str(path),
                }
                for score in SCORE_COLUMNS:
                    row.setdefault(score, pd.NA)
                rows.append(row)
    finally:
        workbook.close()
    return rows


def _visit_header_from_row(values: tuple[Any, ...]) -> tuple[bool, str]:
    for value in values:
        if not isinstance(value, str):
            continue
        match = VISIT_PATTERN.search(value)
        if match:
            visit = f"{int(match.group(1))}w"
            return True, visit if visit in VISIT_PRIORITY else ""
    return False, ""


def _score_column_map(values: tuple[Any, ...]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for index, value in enumerate(values):
        label = str(value).strip() if value is not None else ""
        canonical = _canonical_score(label)
        if canonical and canonical not in mapping:
            mapping[canonical] = index
    return mapping


def _canonical_score(label: str) -> str:
    if label == "嗜睡":
        return "sleepiness"
    upper = label.upper()
    return upper if upper in {"ISI", "PSQI", "BAI", "BDI"} else ""


def _numeric_value(value: Any) -> float | int | Any:
    numeric = pd.to_numeric(value, errors="coerce")
    return numeric if not pd.isna(numeric) else pd.NA


def _subject_from_path(path: Path) -> str:
    for part in reversed(path.parts):
        if re.fullmatch(r"sub-[A-Za-z0-9]+", part):
            return part
    return ""


def _combine_duplicate_visits(frame: pd.DataFrame) -> tuple[pd.DataFrame, int]:
    if frame.empty:
        return pd.DataFrame(columns=LONG_COLUMNS), 0
    duplicate_count = int(frame.duplicated(["subject_id", "visit"]).sum())
    combined: list[dict[str, Any]] = []
    for (_, _), group in frame.groupby(["subject_id", "visit"], sort=False, dropna=False):
        first = group.iloc[0]
        row = {
            "subject_id": first["subject_id"],
            "subject": first["subject"],
            "visit": first["visit"],
            "source_file": first["source_file"],
        }
        for score in SCORE_COLUMNS:
            available = group[score].dropna()
            row[score] = available.iloc[0] if not available.empty else pd.NA
        combined.append(row)
    return pd.DataFrame(combined, columns=LONG_COLUMNS), duplicate_count
